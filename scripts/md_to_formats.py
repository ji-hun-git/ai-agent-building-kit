"""Convert the exported doc tabs (Markdown) to HTML, PDF and DOCX.

The doc export replaces drawings and uploaded images with text placeholders;
this script swaps them back for the PNGs kept in the kit.

Usage: python md_to_formats.py <doc_dir> <images_dir> <chrome_exe>
"""

import base64
import re
import subprocess
import sys
from pathlib import Path

import markdown
from docx import Document
from docx.shared import Inches, Pt

PLACEHOLDER = re.compile(r"&#91;(?:embedded content|image): (.*?)\\\]")
# placeholder caption keyword -> image file in images_dir
IMAGES = {
    "request plane": "diagrams/ontology_request_and_change_planes.png",
    "agent stack": "diagrams/agent_stack_2026.png",
    "Key frames": "scenes_contact_sheet.png",
}
CSS = """
body { font-family: 'Malgun Gothic', 'Segoe UI', sans-serif; max-width: 820px; margin: 32px auto;
       color: #111; line-height: 1.5; font-size: 14px; padding: 0 16px; }
h1 { font-size: 26px; } h2 { font-size: 19px; margin-top: 32px; border-bottom: 1px solid #ccc; }
table { border-collapse: collapse; width: 100%; margin: 12px 0; font-size: 12.5px; }
th, td { border: 1px solid #bbb; padding: 4px 6px; vertical-align: top; text-align: left; }
th { background: #f0f0f0; } code { background: #f4f4f4; padding: 0 3px; }
img { max-width: 100%; border: 1px solid #ddd; margin: 8px 0; }
@media print { body { margin: 0 auto; } h2 { break-after: avoid; } tr { break-inside: avoid; } }
"""


def image_for(caption, images_dir):
    for key, rel in IMAGES.items():
        if key.lower() in caption.lower():
            path = images_dir / rel
            return path if path.exists() else None
    return None


def to_html(md_text, images_dir, title):
    def swap(m):
        img = image_for(m.group(1), images_dir)
        if not img:
            return f"*[{m.group(1)}]*"
        data = base64.b64encode(img.read_bytes()).decode()
        return f'<img alt="{m.group(1)}" src="data:image/png;base64,{data}">'

    body = markdown.markdown(PLACEHOLDER.sub(swap, md_text), extensions=["tables", "sane_lists"])
    return f"<!doctype html><html><head><meta charset='utf-8'><title>{title}</title><style>{CSS}</style></head><body>{body}</body></html>"


def add_runs(par, text):
    """Bold (**x**), code (`x`) and links ([t](u)) -> plain runs, link as 't (u)'."""
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1 (\2)", text)
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            par.add_run(part[2:-2]).bold = True
        elif part.startswith("`"):
            run = par.add_run(part[1:-1])
            run.font.name = "Consolas"
        else:
            par.add_run(part.replace("\\[", "[").replace("\\]", "]"))


def to_docx(md_text, images_dir, out_path):
    doc = Document()
    doc.styles["Normal"].font.name = "Malgun Gothic"
    doc.styles["Normal"].font.size = Pt(10)
    lines = md_text.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        m = PLACEHOLDER.search(line)
        if m:
            img = image_for(m.group(1), images_dir)
            if img:
                doc.add_picture(str(img), width=Inches(6.3))
            else:
                doc.add_paragraph(f"[{m.group(1)}]")
        elif line.startswith("|") and i + 1 < len(lines) and set(lines[i + 1].replace("|", "").strip()) <= set("-: "):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                if not set(lines[i].replace("|", "").strip()) <= set("-: "):
                    rows.append([c.strip() for c in lines[i].strip().strip("|").split(" | ")])
                i += 1
            width = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.style = "Table Grid"
            for r, cells in enumerate(rows):
                for c, cell in enumerate(cells):
                    par = table.cell(r, c).paragraphs[0]
                    add_runs(par, cell)
                    if r == 0:
                        for run in par.runs:
                            run.bold = True
            continue
        elif line.startswith("#"):
            level = len(line) - len(line.lstrip("#"))
            doc.add_heading(line.lstrip("#").strip(), level=min(level, 3))
        elif re.match(r"^\s*([-*]|\d+\.)\s", line):
            style = "List Number" if re.match(r"^\s*\d+\.", line) else "List Bullet"
            add_runs(doc.add_paragraph(style=style), re.sub(r"^\s*([-*]|\d+\.)\s", "", line))
        elif line.strip():
            add_runs(doc.add_paragraph(), line)
        i += 1
    doc.save(out_path)


def main(doc_dir, images_dir, chrome):
    doc_dir, images_dir = Path(doc_dir), Path(images_dir)
    for md in sorted(doc_dir.glob("*.md")):
        text = md.read_text(encoding="utf-8")
        html = md.with_suffix(".html")
        html.write_text(to_html(text, images_dir, md.stem), encoding="utf-8")
        pdf = md.with_suffix(".pdf")
        subprocess.run([chrome, "--headless=new", "--disable-gpu", "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf}", html.resolve().as_uri()], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=120)
        to_docx(text, images_dir, md.with_suffix(".docx"))
        print(md.name, "->", html.name, pdf.name, md.with_suffix(".docx").name)


if __name__ == "__main__":
    main(*sys.argv[1:4])
